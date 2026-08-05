"""Per-job tailored CV generation (2026-07-25 spec).

Applicant-side module — the scrape pipeline must NEVER import it
(tests/test_cowork_import_guard.py). Operates only on COPIES of the CV
docx masters; masters are never written.

Never-fabricate, by construction: content comes exclusively from the CV
itself plus owner-attested keywords (AttestedSkill rows, human-entered).
This module decides placement and order, never content.
"""

from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path

from docx import Document

SEP = " · "                     # the CV's list separator
_EM_DASHES = ("—", "–")  # — and – : owner rule, never generated

# Section headings that delimit calibration parsing (lowercased, exact).
_HEADINGS = {
    "professional profile", "how i build with ai", "professional experience",
    "key projects", "technical skills", "education", "languages",
}

# A project header: "Name   (Stack · Stack · Stack)" — full line, no bullets.
_PROJECT_RE = re.compile(r"^(?P<name>[^()]+?)\s*\((?P<stack>[^()]+)\)$")


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def slug(text: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "_", text or "").strip("_")
    return s or "CV"


def split_items(body: str) -> list[str]:
    return [s.strip() for s in body.split("·") if s.strip()]


def order_for_jd(items: list[str], jd: set[str]) -> list[str]:
    hits = [s for s in items if s.lower() in jd]
    rest = [s for s in items if s.lower() not in jd]
    return hits + rest


def merged_skills_body(
    body: str, attested: list[dict], label: str, jd: set[str]
) -> str:
    """One skills-category line's new body: attested keywords whose
    category_target is `label` inserted (display form), then JD-relevant
    items moved to the front (stable otherwise)."""
    items = split_items(body)
    have = {s.lower() for s in items}
    for a in attested:
        if a.get("category_target") == label and a["keyword_norm"] not in have:
            items.append(a.get("display") or a["keyword_norm"])
            have.add(a["keyword_norm"])
    return SEP.join(order_for_jd(items, jd))


def stack_addition(
    project_name: str, attested: list[dict], jd: set[str], existing_stack: str = "",
) -> list[str]:
    """Attested keywords placed in this project AND asked for by the JD —
    stacks must not bloat on jobs that don't care. Also deduped
    case-insensitively against `existing_stack` (the project's calibrated
    parenthesized stack): a keyword already listed there must never be
    appended again, even if it's missing from the PDF-extracted CV skills
    the owner separately attested it against."""
    have = {s.lower() for s in split_items(existing_stack)}
    return [
        (a.get("display") or a["keyword_norm"])
        for a in attested
        if project_name in (a.get("project_targets") or [])
        and a["keyword_norm"] in jd
        and a["keyword_norm"] not in have
    ]


def render_summary(template: str, top_skills: list[str]) -> str:
    out = template.replace("{skills}", ", ".join(top_skills))
    if any(ch in out for ch in _EM_DASHES):
        raise ValueError("em/en dash in generated summary (owner rule)")
    return out


# JD keywords are normalized lowercase; acronyms and proper nouns must not
# reach an employer-facing summary as "ai" or "javascript". Keywords absent
# here pass through unchanged (generic vocabulary reads fine lowercase).
_DISPLAY_CASE = {
    "ai": "AI", "ml": "ML", "nlp": "NLP", "llm": "LLM", "rag": "RAG",
    "ci-cd": "CI/CD", "api": "API", "rest": "REST", "graphql": "GraphQL",
    "sql": "SQL", "nosql": "NoSQL", "aws": "AWS", "gcp": "GCP",
    "azure": "Azure", "python": "Python", "java": "Java",
    "javascript": "JavaScript", "typescript": "TypeScript", "c#": "C#",
    "c++": "C++", "go": "Go", "rust": "Rust", "php": "PHP", "ruby": "Ruby",
    "kotlin": "Kotlin", "swift": "Swift", "r": "R", "react": "React",
    "angular": "Angular", "vue": "Vue", "nextjs": "Next.js",
    "node": "Node.js", "nodejs": "Node.js", "django": "Django",
    "flask": "Flask", "fastapi": "FastAPI", "spring": "Spring",
    "rails": "Rails", "laravel": "Laravel", "docker": "Docker",
    "kubernetes": "Kubernetes", "terraform": "Terraform", "linux": "Linux",
    "git": "Git", "github": "GitHub", "gitlab": "GitLab",
    "postgres": "PostgreSQL", "postgresql": "PostgreSQL", "mysql": "MySQL",
    "mongodb": "MongoDB", "redis": "Redis", "sqlite": "SQLite",
    "elasticsearch": "Elasticsearch", "kafka": "Kafka", "spark": "Spark",
    "pytorch": "PyTorch", "tensorflow": "TensorFlow", "numpy": "NumPy",
    "pandas": "pandas", "devops": "DevOps", "qa": "QA", "seo": "SEO",
    "crm": "CRM", "sap": "SAP", "excel": "Excel", "salesforce": "Salesforce",
    "english": "English", "arabic": "Arabic", "german": "German",
    "french": "French",
}


def top_skills_for(
    jd_keywords: list[str], matched: list[str], attested: list[dict],
    limit: int = 3,
) -> list[str]:
    """The summary slot: first `limit` JD keywords the applicant actually
    has (CV-matched or attested). Attested display casing wins; CV-matched
    keywords fall back to _DISPLAY_CASE, then pass through as-is."""
    display = {a["keyword_norm"]: (a.get("display") or a["keyword_norm"])
               for a in attested}
    have = {m.lower() for m in matched} | set(display)
    out: list[str] = []
    for kw in jd_keywords:
        if kw in have:
            out.append(display.get(kw, _DISPLAY_CASE.get(kw, kw)))
        if len(out) == limit:
            break
    return out


def calibrate_docx(path: str | Path) -> dict:
    """Parse a CV master into exact-string anchors. Fails loudly (ValueError
    naming what's missing) rather than guessing — the /profile calibration
    surface shows the owner exactly what was detected for confirmation."""
    doc = Document(str(path))
    current = ""
    summary = ""
    skills_lines: list[dict] = []
    projects: list[dict] = []
    for p in doc.paragraphs:
        stripped = p.text.strip()
        if not stripped:
            continue
        if stripped.lower() in _HEADINGS:
            current = stripped.lower()
            continue
        if current == "professional profile" and not summary:
            summary = p.text
        elif current == "technical skills" and ":" in stripped:
            skills_lines.append(
                {"label": stripped.split(":", 1)[0].strip(), "text": p.text}
            )
        elif current == "key projects" and not stripped.startswith("•"):
            m = _PROJECT_RE.match(stripped)
            if m:
                projects.append({
                    "name": m.group("name").strip(),
                    "stack": m.group("stack").strip(),
                    "text": p.text,
                })
    missing = [
        name for name, ok in (
            ("summary", summary),
            ("skills lines", skills_lines),
            ("projects", projects),
        ) if not ok
    ]
    if missing:
        raise ValueError(f"calibration failed: no {', '.join(missing)} found")
    return {
        "sha256": file_sha256(path),
        "summary": summary,
        "skills_lines": skills_lines,
        "projects": projects,
    }


def _find_paragraph(doc, anchor_text: str):
    for p in doc.paragraphs:
        if p.text == anchor_text:
            return p
    return None


def _label_prefix(anchor_text: str) -> str:
    """'Backend / Frontend:   FastAPI · …' → 'Backend / Frontend:   '."""
    i = anchor_text.index(":") + 1
    while i < len(anchor_text) and anchor_text[i] == " ":
        i += 1
    return anchor_text[:i]


def _replace_body_after_label(p, prefix: str, new_body: str) -> bool:
    """Rewrite everything after `prefix`, editing runs in place so the
    label's formatting is untouched. False when the prefix boundary can't
    be located cleanly — caller skips + flags, never guesses."""
    if not p.text.startswith(prefix) or not p.runs:
        return False
    pos = 0
    for i, r in enumerate(p.runs):
        end = pos + len(r.text)
        if end >= len(prefix):
            keep = prefix[pos:]
            if not r.text.startswith(keep):
                return False
            r.text = keep + new_body
            for later in p.runs[i + 1:]:
                later.text = ""
            return True
        pos = end
    return False


def _append_into_stack(p, additions: list[str]) -> bool:
    """Insert ' · X · Y' before the stack's closing paren, in whichever run
    holds it — formatting elsewhere untouched."""
    if not additions:
        return True
    for r in reversed(p.runs):
        idx = r.text.rfind(")")
        if idx != -1:
            r.text = r.text[:idx] + SEP + SEP.join(additions) + r.text[idx:]
            return True
    return False


def _set_paragraph_text(p, new_text: str) -> bool:
    """Whole-paragraph replace (summary only): first run keeps its
    formatting and takes the text; the rest are cleared."""
    if not p.runs:
        return False
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""
    return True


def _apply_tailoring_edits(
    out_path: Path,
    anchors: dict,
    attested: list[dict],
    jd_keywords: list[str],
    matched: list[str],
    summary_template: str,
) -> tuple[bool, str]:
    """The edit body, operating on the already-copied `out_path`. Returns
    (fully_tailored, reason) — never touches the master, never decides
    whether to keep or remove `out_path` (the caller owns that)."""
    doc = Document(str(out_path))
    jd = {k.lower() for k in jd_keywords}
    skipped: list[str] = []

    for line in anchors.get("skills_lines", []):
        p = _find_paragraph(doc, line["text"])
        if p is None or ":" not in line["text"]:
            skipped.append(f"skills line '{line['label']}'")
            continue
        prefix = _label_prefix(line["text"])
        body = line["text"][len(prefix):]
        new_body = merged_skills_body(body, attested, line["label"], jd)
        if not _replace_body_after_label(p, prefix, new_body):
            skipped.append(f"skills line '{line['label']}'")

    for proj in anchors.get("projects", []):
        additions = stack_addition(proj["name"], attested, jd, proj.get("stack", ""))
        if not additions:
            continue
        p = _find_paragraph(doc, proj["text"])
        if p is None or not _append_into_stack(p, additions):
            skipped.append(f"project '{proj['name']}'")

    if summary_template.strip():
        top = top_skills_for(jd_keywords, matched, attested)
        p = _find_paragraph(doc, anchors.get("summary", ""))
        if p is None or not top:
            skipped.append("summary")
        else:
            try:
                if not _set_paragraph_text(p, render_summary(summary_template, top)):
                    skipped.append("summary")
            except ValueError:
                return False, "summary template renders an em dash (owner rule)"

    doc.save(str(out_path))
    if skipped:
        return False, "partial tailoring; skipped: " + ", ".join(skipped)
    return True, "tailored from attested skills"


def generate_tailored_docx(
    master_path: str,
    anchors: dict,
    attested: list[dict],
    jd_keywords: list[str],
    matched: list[str],
    summary_template: str,
    out_path: Path,
) -> tuple[bool, str]:
    """Copy the master to out_path and apply the three deterministic edits
    (skills lines, project stacks, summary). Returns (fully_tailored,
    reason); the caller treats anything but (True, …) as fall-back-to-PDF.
    The master itself is never written. Anything other than (True, …) after
    the copy was made removes `out_path` again — the filename is stable per
    application, so leaving a pristine or partially-edited copy behind
    would silently overwrite (or masquerade as) a previously good tailored
    file the next time this application is exported."""
    if file_sha256(master_path) != anchors.get("sha256"):
        return False, "cv changed since calibration; recalibrate on /profile"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(master_path, out_path)
    ok, reason = _apply_tailoring_edits(
        out_path, anchors, attested, jd_keywords, matched, summary_template
    )
    if not ok:
        out_path.unlink(missing_ok=True)
    return ok, reason
