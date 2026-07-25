"""cv_docx: pure tailoring logic (Task 2) + calibration/generation (Tasks 3-4)."""
from __future__ import annotations

import pytest

ATTESTED = [
    {"keyword_norm": "graphql", "display": "GraphQL",
     "category_target": "Backend / Frontend",
     "project_targets": ["UniVeranstaltungen - Events Platform"]},
    {"keyword_norm": "security", "display": "Security",
     "category_target": "Data / DevOps", "project_targets": []},
]


def test_order_for_jd_stable():
    from jobhunt.cv_docx import order_for_jd
    items = ["LangChain", "RAG", "FAISS"]
    assert order_for_jd(items, {"faiss", "rag"}) == ["RAG", "FAISS", "LangChain"]
    assert order_for_jd(items, set()) == items


def test_merged_skills_body_inserts_attested_and_leads_jd():
    from jobhunt.cv_docx import merged_skills_body
    body = "FastAPI · REST · React"
    out = merged_skills_body(body, ATTESTED, "Backend / Frontend", {"graphql", "react"})
    # attested GraphQL inserted; JD-relevant (react, graphql) lead, original order otherwise
    assert out == "React · GraphQL · FastAPI · REST"


def test_merged_skills_body_no_duplicate_insert():
    from jobhunt.cv_docx import merged_skills_body
    out = merged_skills_body("GraphQL · REST", ATTESTED, "Backend / Frontend", set())
    assert out.count("GraphQL") == 1


def test_stack_addition_only_when_jd_asks():
    from jobhunt.cv_docx import stack_addition
    assert stack_addition(
        "UniVeranstaltungen - Events Platform", ATTESTED, {"graphql"}
    ) == ["GraphQL"]
    assert stack_addition(
        "UniVeranstaltungen - Events Platform", ATTESTED, {"python"}
    ) == []
    assert stack_addition(
        "DocAI - Intelligent Document Processing", ATTESTED, {"graphql"}
    ) == []


def test_stack_addition_dedupes_against_existing_stack():
    from jobhunt.cv_docx import stack_addition
    # "Node" is already in the fixture project's stack
    # (React · TypeScript · Node) -- an attestation for it must not
    # produce a duplicate "(React · TypeScript · Node · Node)".
    attested_node = [{"keyword_norm": "node", "display": "Node",
                       "category_target": None,
                       "project_targets": ["UniVeranstaltungen - Events Platform"]}]
    assert stack_addition(
        "UniVeranstaltungen - Events Platform", attested_node, {"node"},
        "React · TypeScript · Node",
    ) == []
    # a genuinely new keyword still gets added
    assert stack_addition(
        "UniVeranstaltungen - Events Platform", attested_node, {"node"}, "React",
    ) == ["Node"]


def test_render_summary_and_em_dash_guard():
    from jobhunt.cv_docx import render_summary
    assert render_summary("Engineer with {skills}.", ["RAG", "GraphQL"]) == \
        "Engineer with RAG, GraphQL."
    with pytest.raises(ValueError):
        render_summary("Engineer — with {skills}.", ["RAG"])


def test_top_skills_for_jd_ordered_max3():
    from jobhunt.cv_docx import top_skills_for
    jd = ["ai", "graphql", "python", "rag", "rest"]
    out = top_skills_for(jd, matched=["ai", "python", "rag", "rest"], attested=ATTESTED)
    assert out == ["ai", "GraphQL", "python"]  # first 3 in JD order, display-cased


def test_slug():
    from jobhunt.cv_docx import slug
    assert slug("GitLab, Inc.") == "GitLab_Inc"
    assert slug("") == "CV"


def make_cv_docx(path):
    """Fixture mirroring the real masters' layout, incl. multi-run
    paragraphs (bold labels) so run-level editing is genuinely exercised."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Ziad Ahmed Abdalla")
    doc.add_paragraph("Professional Profile")
    doc.add_paragraph("AI engineer who ships products end-to-end.")
    doc.add_paragraph("Key Projects")
    proj = doc.add_paragraph()
    proj.add_run("UniVeranstaltungen - Events Platform   ").bold = True
    proj.add_run("(React · TypeScript · Node)")
    doc.add_paragraph("•  100+ components (JWT RBAC) shipped")
    doc.add_paragraph("Technical Skills")
    line = doc.add_paragraph()
    line.add_run("Backend / Frontend:").bold = True
    line.add_run("   FastAPI · REST · React")
    doc.add_paragraph("Data / DevOps:   PostgreSQL · Docker")
    doc.add_paragraph("Education")
    doc.add_paragraph("German University in Cairo")
    doc.save(str(path))


def test_calibrate_docx(tmp_path):
    from jobhunt.cv_docx import calibrate_docx, file_sha256
    f = tmp_path / "cv.docx"
    make_cv_docx(f)
    cal = calibrate_docx(f)
    assert cal["sha256"] == file_sha256(f)
    assert cal["summary"] == "AI engineer who ships products end-to-end."
    assert [line["label"] for line in cal["skills_lines"]] == \
        ["Backend / Frontend", "Data / DevOps"]
    assert cal["projects"][0]["name"] == "UniVeranstaltungen - Events Platform"
    assert cal["projects"][0]["stack"] == "React · TypeScript · Node"
    # the bullet with parentheses must NOT be a project
    assert len(cal["projects"]) == 1


def test_calibrate_docx_missing_sections_raises(tmp_path):
    from docx import Document

    from jobhunt.cv_docx import calibrate_docx

    f = tmp_path / "bad.docx"
    d = Document()
    d.add_paragraph("Just a name")
    d.save(str(f))
    with pytest.raises(ValueError, match="summary"):
        calibrate_docx(f)


def _gen(tmp_path, attested, jd, matched, template=""):
    from jobhunt.cv_docx import calibrate_docx, generate_tailored_docx
    master = tmp_path / "master.docx"
    make_cv_docx(master)
    cal = calibrate_docx(master)
    out = tmp_path / "out" / "T_CV_X.docx"
    ok, reason = generate_tailored_docx(
        str(master), cal, attested, jd, matched, template, out)
    return master, out, ok, reason


def _texts(path):
    from docx import Document
    return [p.text for p in Document(str(path)).paragraphs]


def test_generate_edits_skills_projects_summary(tmp_path):
    master, out, ok, reason = _gen(
        tmp_path, ATTESTED, ["graphql", "react"], ["react"],
        template="Engineer with {skills}.")
    assert ok, reason
    texts = _texts(out)
    assert "Backend / Frontend:   React · GraphQL · FastAPI · REST" in texts
    assert any("(React · TypeScript · Node · GraphQL)" in t for t in texts)
    # jd order is ["graphql", "react"] → GraphQL (attested display) leads
    assert "Engineer with GraphQL, react." in texts
    # master untouched
    assert "Backend / Frontend:   FastAPI · REST · React" in _texts(master)


def test_generate_no_duplicate_stack_keyword(tmp_path):
    """An attested keyword already in the docx stack (e.g. 'Node' parsed
    from the fixture project's stack, but missing from the PDF-extracted
    CV skills so the owner attested it anyway) must not be duplicated."""
    attested = [{"keyword_norm": "node", "display": "Node",
                 "category_target": None,
                 "project_targets": ["UniVeranstaltungen - Events Platform"]}]
    master, out, ok, reason = _gen(tmp_path, attested, ["node"], [])
    assert ok, reason
    texts = _texts(out)
    stack_line = next(t for t in texts if "UniVeranstaltungen" in t)
    assert stack_line.count("Node") == 1
    assert "(React · TypeScript · Node)" in stack_line


def test_generate_preserves_label_formatting(tmp_path):
    from docx import Document
    _, out, ok, _ = _gen(tmp_path, ATTESTED, ["graphql"], [])
    assert ok
    doc = Document(str(out))
    line = next(p for p in doc.paragraphs
                if p.text.startswith("Backend / Frontend:"))
    assert line.runs[0].bold is True
    assert line.runs[0].text == "Backend / Frontend:"


def test_generate_drift_aborts(tmp_path):
    from jobhunt.cv_docx import calibrate_docx, generate_tailored_docx
    master = tmp_path / "m.docx"
    make_cv_docx(master)
    cal = calibrate_docx(master)
    master.write_bytes(master.read_bytes() + b"x")  # master changed post-calibration
    out = tmp_path / "o.docx"
    ok, reason = generate_tailored_docx(str(master), cal, [], [], [], "", out)
    assert not ok and "recalibrate" in reason
    assert not out.exists()  # drift aborts BEFORE copying


def test_generate_anchor_miss_is_partial_not_crash(tmp_path):
    from jobhunt.cv_docx import calibrate_docx, file_sha256, generate_tailored_docx
    master = tmp_path / "m.docx"
    make_cv_docx(master)
    cal = calibrate_docx(master)
    cal["skills_lines"][0]["text"] = "NOT IN THE DOCUMENT"
    cal["sha256"] = file_sha256(master)
    out = tmp_path / "o.docx"
    ok, reason = generate_tailored_docx(
        str(master), cal, ATTESTED, ["graphql"], [], "", out)
    assert not ok and "skipped" in reason and "Backend / Frontend" in reason
    # a partial (non-True) result must not leave a stray/stale docx behind —
    # the filename is stable per application, so a leftover partial copy
    # could masquerade as (or overwrite) a previously good tailored file.
    assert not out.exists()


def test_generate_em_dash_summary_leaves_no_stray_file(tmp_path):
    """The em-dash guard path (a caller-supplied summary_template that
    would render a dash) also must not leave a partially-edited docx
    behind — same stray-file hazard as the anchor-miss path."""
    master, out, ok, reason = _gen(
        tmp_path, ATTESTED, ["graphql"], ["graphql"],
        template="Engineer — {skills}.")  # em dash
    assert not ok and "em dash" in reason
    assert not out.exists()
